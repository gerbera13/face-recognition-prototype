"""Конвертация markdown-отчётов в .docx (Word)"""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os, re

BASE = os.path.dirname(__file__)

def md_to_docx(md_path, docx_path, title):
    doc = Document()

    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    for line in lines:
        line = line.rstrip()

        if line.startswith('# ') and not line.startswith('## '):
            h = doc.add_heading(line[2:], level=0)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=1)
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=2)
        elif line.startswith('---'):
            doc.add_paragraph('_' * 50)
        elif line.startswith('- **'):
            p = doc.add_paragraph()
            r = p.add_run(line[2:line.index('**:')])
            r.bold = True
            p.add_run(line[line.index('**:')+3:])
        elif line.startswith('- '):
            doc.add_paragraph(line[2:], style='List Bullet')
        elif line.startswith('|'):
            continue  # skip table rows, they look bad in simple conversion
        elif line.strip() == '':
            continue
        elif re.match(r'^\d+\.\s', line):
            doc.add_paragraph(re.sub(r'^\d+\.\s', '', line), style='List Number')
        else:
            # bold markers
            p = doc.add_paragraph()
            parts = re.split(r'(\*\*.*?\*\*)', line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                else:
                    p.add_run(part)

    doc.save(docx_path)
    print(f'  [+] {os.path.basename(docx_path)}')


md_to_docx(os.path.join(BASE, 'report.md'), os.path.join(BASE, 'Отчёт.docx'), 'Отчёт')
md_to_docx(os.path.join(BASE, 'presentation.md'), os.path.join(BASE, 'Презентация.docx'), 'Презентация')
md_to_docx(os.path.join(BASE, 'questions.md'), os.path.join(BASE, 'Вопросы.docx'), 'Вопросы')
print('\nГотово.')
